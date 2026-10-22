"""Builders for small sample/hostile files used by tests and the demo."""
from __future__ import annotations

import io
import zipfile


def pdf_bytes(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        if text:
            c.setFont("HeiseiKakuGo-W5", 12)
            y = 800
            for line in text.split("\n"):
                c.drawString(50, y, line)
                y -= 18
        else:
            c.rect(100, 400, 200, 200, fill=1)  # image-like page without text layer
        c.showPage()
    c.save()
    return buf.getvalue()


def docx_bytes(paras: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("仕様書", level=1)
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def xlsx_bytes(rows: list[list], sheet: str = "単価表", formula_cell: bool = True) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    if formula_cell:
        ws.append(["合計", "=SUM(B2:B3)"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pptx_bytes(slides: list[str]) -> bytes:
    from pptx import Presentation
    prs = Presentation()
    for s in slides:
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = s.split("\n")[0]
        sl.placeholders[1].text = "\n".join(s.split("\n")[1:]) or s
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def png_bytes(w: int = 4, h: int = 4) -> bytes:
    import struct
    import zlib
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def tamper_zip(original: bytes, add: dict[str, bytes] | None = None, replace: dict[str, bytes] | None = None) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(original))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for i in src.infolist():
            data = (replace or {}).get(i.filename, src.read(i.filename))
            z.writestr(i, data)
        for n, d in (add or {}).items():
            z.writestr(n, d)
    return buf.getvalue()


def slow_worker(path, kind, limits, out_path):
    """Spawned-process stand-in that never finishes in time (timeout test)."""
    import time
    time.sleep(30)
