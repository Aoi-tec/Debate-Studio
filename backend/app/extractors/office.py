"""docx / xlsx / pptx extractors (python-docx, openpyxl, python-pptx).

Formulas are never recalculated; cached values may be missing or stale.
"""
from __future__ import annotations

from pathlib import Path

from app.extractors.base import Chunk, ExtractionError, ExtractResult, Limits, split_lines_into_chunks
from app.core.tokens import estimate_tokens


def extract_docx(path: Path, limits: Limits) -> ExtractResult:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        d = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        raise ExtractionError("PARSE_ERROR", f"Word文書を解析できません: {type(e).__name__}") from e
    res = ExtractResult(stats={"units_total": 1})
    blocks: list[tuple[str, dict]] = []
    heading = ""
    p_index = 0
    t_index = 0
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p_index += 1
            para = Paragraph(child, d)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith("heading") or style.startswith("見出し"):
                heading = text
                text = f"## {text}"
            blocks.append((text, {"paragraph": p_index, "heading": heading}))
        elif tag == "tbl":
            t_index += 1
            table = Table(child, d)
            rows = []
            for r in table.rows:
                cells = []
                for c in r.cells:
                    cells.append(c.text.strip().replace("\n", " "))
                rows.append(" | ".join(cells))
            if rows:
                blocks.append(("[表]\n" + "\n".join(rows), {"table": t_index, "heading": heading}))
    # group blocks into chunks with paragraph ranges
    buf: list[str] = []
    locs: list[dict] = []
    tok = 0

    def flush() -> None:
        nonlocal buf, locs, tok
        if not buf:
            return
        loc: dict = {}
        paras = [x["paragraph"] for x in locs if "paragraph" in x]
        tables = [x["table"] for x in locs if "table" in x]
        if paras:
            loc["paragraph_start"], loc["paragraph_end"] = min(paras), max(paras)
        if tables:
            loc["tables"] = tables
        if locs and locs[0].get("heading"):
            loc["heading"] = locs[0]["heading"]
        res.chunks.append(Chunk("\n".join(buf), loc))
        buf, locs, tok = [], [], 0

    for text, loc in blocks:
        t = estimate_tokens(text)
        if buf and tok + t > 400:
            flush()
        buf.append(text)
        locs.append(loc)
        tok += t
    flush()
    try:
        n_img = len(d.inline_shapes)
    except Exception:  # noqa: BLE001
        n_img = 0
    if n_img:
        res.unreadable.append({"location": {}, "reason": f"埋め込み画像{n_img}件の内容は未解析"})
    res.stats.update(paragraphs=p_index, tables=t_index)
    return res


def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def extract_xlsx(path: Path, limits: Limits) -> ExtractResult:
    import openpyxl

    try:  # file objects: stored names intentionally have no extension
        fh_f = open(path, "rb")
        fh_v = open(path, "rb")
        wb_f = openpyxl.load_workbook(fh_f, read_only=True, data_only=False)
        wb_v = openpyxl.load_workbook(fh_v, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ExtractionError("PARSE_ERROR", f"Excelを解析できません: {type(e).__name__}") from e
    res = ExtractResult(stats={"sheets": len(wb_f.sheetnames), "units_total": len(wb_f.sheetnames)})
    cells_seen = 0
    missing_cache = 0
    for name in wb_f.sheetnames:
        ws_f = wb_f[name]
        ws_v = wb_v[name]
        if not hasattr(ws_f, "iter_rows"):
            res.unreadable.append({"location": {"sheet": name}, "reason": "グラフシートは未解析"})
            continue
        lines: list[tuple[int, str]] = []
        max_col = 0
        for row_f, row_v in zip(ws_f.iter_rows(), ws_v.iter_rows()):
            parts = []
            r_idx = None
            for cf, cv in zip(row_f, row_v):
                if getattr(cf, "value", None) is None:
                    continue
                cells_seen += 1
                if cells_seen > limits.max_xlsx_cells:
                    raise ExtractionError("TOO_MANY_CELLS", f"セル数が上限{limits.max_xlsx_cells}を超えています")
                r_idx = cf.row
                max_col = max(max_col, cf.column)
                coord = f"{_col(cf.column)}{cf.row}"
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    cached = cv.value
                    if cached is None:
                        missing_cache += 1
                        parts.append(f"{coord}={cf.value}（計算値なし）")
                    else:
                        parts.append(f"{coord}={cached}（式 {cf.value}・キャッシュ値）")
                else:
                    parts.append(f"{coord}={cf.value}")
            if parts and r_idx is not None:
                lines.append((r_idx, " | ".join(parts)))
        if not lines:
            continue
        # chunk by rows
        buf: list[tuple[int, str]] = []
        tok = 0
        for r, line in lines + [(None, None)]:  # type: ignore[list-item]
            if line is None or (buf and tok + estimate_tokens(line) > 400):
                if buf:
                    rng = f"A{buf[0][0]}:{_col(max(max_col, 1))}{buf[-1][0]}"
                    res.chunks.append(Chunk(f"[シート {name}]\n" + "\n".join(x[1] for x in buf),
                                            {"sheet": name, "cell_range": rng}))
                buf, tok = [], 0
                if line is None:
                    break
            buf.append((r, line))
            tok += estimate_tokens(line)
    if missing_cache:
        res.unreadable.append({"location": {}, "reason": f"数式{missing_cache}件はキャッシュ済み計算値がなく値不明（再計算はしない）"})
    res.stats["cells"] = cells_seen
    wb_f.close()
    wb_v.close()
    fh_f.close()
    fh_v.close()
    return res


def extract_pptx(path: Path, limits: Limits) -> ExtractResult:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        raise ExtractionError("PARSE_ERROR", f"PowerPointを解析できません: {type(e).__name__}") from e
    slides = list(prs.slides)
    res = ExtractResult(stats={"slides": len(slides), "units_total": len(slides)})
    pictures = 0
    for i, slide in enumerate(slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [" | ".join(c.text.strip() for c in r.cells) for r in shape.table.rows]
                texts.append("[表]\n" + "\n".join(rows))
            if shape.shape_type == 13:  # PICTURE
                pictures += 1
        if slide.has_notes_slide:
            n = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if n:
                texts.append(f"[ノート] {n}")
        if texts:
            res.chunks.extend(split_lines_into_chunks("\n".join(texts).split("\n"),
                                                      lambda s, e, sl=i: {"slide": sl, "line_start": s, "line_end": e}))
        else:
            res.unreadable.append({"location": {"slide": i}, "reason": "テキストなし（図や画像のみの可能性）"})
    if pictures:
        res.unreadable.append({"location": {}, "reason": f"画像{pictures}件の内容は未解析"})
    return res
